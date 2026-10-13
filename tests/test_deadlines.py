import io
from datetime import datetime

import pytest

from planner.deadlines import find_deadlines
from planner.documents import UnsupportedDocument, extract_text, is_supported
from planner.rules import TZ

POSTED = datetime(2026, 10, 13, 10, 0, tzinfo=TZ)   # Dienstag


def dl(text, **kw):
    return find_deadlines(text, posted=kw.get("posted", POSTED), now=kw.get("now", POSTED))


def at(d, m, hh=23, mm=59, y=2026):
    return datetime(y, m, d, hh, mm, tzinfo=TZ)


# -- die Beispiele des Nutzers -----------------------------------------------------------------------------
def test_explicit_deadline_with_time():
    [d] = dl("Aufgabe 3: Reichen Sie die Lösung als PDF ein.\nAbgabe bis spätestens 20.10. 19 Uhr in ILIAS.")
    assert d.when == at(20, 10, 19, 0) and not d.date_only and not d.past
    assert d.label == "Abgabe" and "20.10." in d.evidence


def test_next_lesson_is_relative():
    [d] = dl("Bitte füllen Sie das neue Dokument bis zur nächsten Stunde aus.")
    assert d.relative == "next_session" and d.when is None


@pytest.mark.parametrize("text", [
    "Abgabe bis spätestens 20.10.2026 um 19:00 Uhr",
    "Abgabe: 20. Oktober 2026, 19 Uhr",
    "Deadline 2026-10-20 19:00",
    "Abgabe bis spätestens 20.10.26 19.00 Uhr",
    "Submission deadline: Oct 20, 2026 at 7:00 pm",
    "Lösungen bis zum 20.10. 19:00 Uhr abgeben.",
])
def test_many_notations_same_instant(text):
    [d] = dl(text)
    assert d.when == at(20, 10, 19, 0), text


def test_wrapped_sentence_date_on_next_line():
    [d] = dl("Die Lösung ist bis spätestens\n25.10. 18:30 Uhr hochzuladen.")
    assert d.when == at(25, 10, 18, 30)


# -- Tabellen (wie echte Zeitpläne) ------------------------------------------------------------------------------
SCHEDULE = """VL Zeitplan
1 29.09.26
2 06.10.26 Abgabe1: Team (8.10.)
3 13.10.26 Abgabe2: Thema
4 20.10.26 Erste Recherchen   Abgabe 3: Vorgehensplanung
17
19.01.27 Abgabe 4: Businessplan am 19.1.2027
"""


def test_table_rows_get_their_own_dates():
    got = {d.label: d for d in dl(SCHEDULE)}
    assert got["Abgabe1: Team"].when == at(8, 10) and got["Abgabe1: Team"].past          # vor dem Veröffentlichungstag
    assert got["Abgabe2: Thema"].when == at(13, 10)                                       # nicht das Datum der Folgezeile!
    assert got["Abgabe 3: Vorgehensplanung"].when == at(20, 10)
    assert got["Abgabe 4: Businessplan"].when == at(19, 1, y=2027)
    assert len(got) == 4


def test_last_semesters_schedule_is_entirely_past():
    old = SCHEDULE.replace(".26", ".25").replace("2027", "2026").replace(".27", ".26")
    found = dl(old)
    assert found and all(d.past for d in found)         # nichts davon darf in den Kalender


# -- Relatives ----------------------------------------------------------------------------------------------------
def test_tomorrow_and_weekday():
    [m] = dl("Die Übung ist bis morgen 12 Uhr abzugeben.")
    assert m.when == at(14, 10, 12, 0)
    [f] = dl("Hausaufgabe: Kapitel 3 lesen bis Freitag.")
    assert f.when == at(16, 10) and f.date_only
    [w] = dl("Bitte bis Dienstag erledigen.")
    assert w.when == at(20, 10)                          # nächster Dienstag, nicht heute


def test_english_phrases():
    [d] = dl("Please submit your solution by next class.")
    assert d.relative == "next_session"
    [e] = dl("Assignment 2 is due Oct 25, 2026 at 11:59 pm.")
    assert e.when == at(25, 10, 23, 59)


# -- Dinge, die KEINE Fristen sind ----------------------------------------------------------------------------------
@pytest.mark.parametrize("text", [
    "Vorlesung von 9 bis 11 Uhr im Raum R-203.",
    "Prüfungsanmeldung bis 5.11.",                         # keine Aufgabe
    "Version 2.1.3 der Software, Kapitel 1.5",
    "Praktische Arbeit: Die Bearbeitung erfolgt in Gruppen. Aufgabenstellung: eine Idee entwickeln.",
    "Die Vorlesung findet bis zur nächsten Woche nicht statt.",
    "",
])
def test_no_false_deadlines(text):
    assert dl(text) == []


def test_year_inference_wraps_into_next_year():
    posted = datetime(2026, 12, 15, 9, 0, tzinfo=TZ)
    [d] = find_deadlines("Abgabe bis 20.1.", posted=posted, now=posted)
    assert d.when == at(20, 1, y=2027) and not d.past


def test_past_flag_relative_to_now_not_posted():
    later = datetime(2026, 10, 25, 9, 0, tzinfo=TZ)
    [d] = find_deadlines("Abgabe bis 20.10.2026 19 Uhr", posted=POSTED, now=later)
    assert d.past


def test_duplicates_collapse_and_results_are_capped():
    assert len(dl("Abgabe bis 20.10. 19 Uhr.\nAbgabe bis 20.10. 19 Uhr.")) == 1
    many = "\n".join(f"Abgabe {i}: Teil {i} am {i}.11.2026" for i in range(1, 30))
    assert len(dl(many)) == 12


# -- Dokument-Leser ---------------------------------------------------------------------------------------------------
def test_text_and_support():
    assert extract_text(b"Abgabe bis 20.10.", "a.txt") == "Abgabe bis 20.10."
    assert is_supported("Plan.xlsx.pdf") and is_supported("x", "application/pdf") and not is_supported("film.mp4")
    with pytest.raises(UnsupportedDocument):
        extract_text(b"\x00", "film.mp4")


def test_docx_and_xlsx_roundtrip():
    from docx import Document
    from openpyxl import Workbook

    doc = Document()
    doc.add_paragraph("Abgabe bis spätestens 20.10. 19 Uhr")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Abgabe 2", "27.10.2026"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text(buf.getvalue(), "blatt.docx")
    assert "spätestens 20.10." in text and "Abgabe 2  27.10.2026" in text

    wb = Workbook()
    wb.active.append(["Abgabe 1", datetime(2026, 10, 20, 19, 0)])
    wb.active.append(["Abgabe 2", datetime(2026, 10, 27)])
    buf = io.BytesIO()
    wb.save(buf)
    text = extract_text(buf.getvalue(), "plan.xlsx")
    assert "Abgabe 1  20.10.2026 19:00" in text and "Abgabe 2  27.10.2026" in text
    assert [d.when for d in dl(text)] == [at(20, 10, 19, 0), at(27, 10)]   # Excel-Daten sind direkt auswertbar


# -- Fehltreffer aus dem Durchlauf durch echte Kurse (Regression) --------------------------------------------
@pytest.mark.parametrize("text", [
    "Anleihe fällig, Restlaufzeit 3 Jahre und Zinszahlung zum 1.2.",
    "In Ihrem Bestand befindet sich ein Put. Fällig 4.3. Berechnen Sie den Wert.",
])
def test_finance_exam_wording_is_not_a_deadline(text):
    assert dl(text) == []


def test_faellig_with_bis_still_counts():
    [d] = dl("Das Formular ist fällig bis 20.10.2026 12 Uhr.")
    assert d.when == at(20, 10, 12, 0)


def test_year_guessed_flag():
    [guessed] = dl("Abgabe bis 20.10.")
    [explicit] = dl("Abgabe bis 20.10.2026")
    assert guessed.year_guessed and not explicit.year_guessed
