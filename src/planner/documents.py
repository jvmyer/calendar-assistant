"""Text aus heruntergeladenen Dokumenten lesen (PDF, Word, Excel, Text)."""

from __future__ import annotations

import io
import logging
from datetime import date, datetime
from pathlib import PurePosixPath

MAX_CHARS = 80_000
MAX_PDF_PAGES = 80
MAX_XLSX_ROWS = 600
SUPPORTED = {".pdf", ".docx", ".xlsx", ".txt", ".md", ".csv"}

logging.getLogger("pypdf").setLevel(logging.ERROR)   # „Ignoring wrong pointing object …“ ist harmlos


class UnsupportedDocument(Exception):
    """Dateityp wird nicht gelesen (Video, Archiv, Bild …)."""


def suffix_of(filename: str, content_type: str = "") -> str:
    suffix = PurePosixPath(filename.lower()).suffix
    if suffix in SUPPORTED:
        return suffix
    # Doppelendungen wie „plan.xlsx.pdf“ zählen nach der letzten Endung – ok. Sonst nach Content-Type:
    return {"application/pdf": ".pdf", "text/plain": ".txt", "text/csv": ".csv",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx"}.get(content_type, suffix)


def is_supported(filename: str, content_type: str = "") -> bool:
    return suffix_of(filename, content_type) in SUPPORTED


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=False)
    return "\n".join((page.extract_text() or "") for page in reader.pages[:MAX_PDF_PAGES])


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("  ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _cell(value) -> str:
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M") if (value.hour or value.minute) else value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return "" if value is None else str(value)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= MAX_XLSX_ROWS:
                break
            cells = [_cell(v) for v in row]
            if any(cells):
                lines.append("  ".join(c for c in cells if c))
    return "\n".join(lines)


def extract_text(data: bytes, filename: str, content_type: str = "") -> str:
    suffix = suffix_of(filename, content_type)
    if suffix == ".pdf":
        text = _pdf(data)
    elif suffix == ".docx":
        text = _docx(data)
    elif suffix == ".xlsx":
        text = _xlsx(data)
    elif suffix in {".txt", ".md", ".csv"}:
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedDocument(suffix or "unbekannter Typ")
    return text[:MAX_CHARS]
